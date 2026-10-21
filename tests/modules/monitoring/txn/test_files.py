"""File ingest: the readers, the sealed upload, the paged job and the error report.

Whenever a test expects an exception from inside a session, ``pytest.raises`` is
the outermost context manager, so the error crosses the session and rolls it back.
"""

import csv
import io
import json
import logging
from collections.abc import Generator, Sequence
from pathlib import Path
from uuid import UUID

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from tests.helpers import audit_events, staff_session, system_ctx
from tests.modules.monitoring.support import ingest_items, make_item
from xolotl.modules.monitoring.enums import BatchSource, FileStatus
from xolotl.modules.monitoring.errors import FileInvalidError, FileTooLargeError
from xolotl.modules.monitoring.events import TRANSACTION_FILE_COMPLETED
from xolotl.modules.monitoring.kinds import INGEST_FILE_KIND
from xolotl.modules.monitoring.txn import files
from xolotl.modules.monitoring.txn.files import (
    CSV_TYPE,
    JSONL_TYPE,
    XLSX_TYPE,
    IngestFileJob,
    RowError,
    count_rows,
    error_report_csv,
    iter_rows,
    read_header,
    read_object,
    row_to_item,
    upload_file,
)
from xolotl.modules.monitoring.txn.ingest import BatchReport, ingest_batch
from xolotl.modules.monitoring.txn.models import (
    Transaction,
    TransactionBatch,
    TransactionFile,
)
from xolotl.modules.monitoring.txn.schema import (
    TransactionInput,
    content_hash,
    validate_batch,
)
from xolotl.modules.party.models import CustomerRelationship
from xolotl.modules.tenants.service import CreatedTenant
from xolotl.platform import clock
from xolotl.platform.config import get_settings
from xolotl.platform.crypto.blobs import open_blob
from xolotl.platform.db import get_sessionmaker, tenant_session
from xolotl.platform.errors import NotFoundError
from xolotl.platform.events.event import event_from_row
from xolotl.platform.events.models import OutboxRow
from xolotl.platform.ids import uuid7
from xolotl.platform.jobs.models import JobRow
from xolotl.platform.jobs.queue import Job, PostgresJobQueue, register_job_handler
from xolotl.platform.jobs.runner import JobRunner
from xolotl.platform.storage import ObjectStore
from xolotl.platform.storage.local import LocalFsObjectStore

FIXTURES = Path(__file__).parent / "fixtures"
DUPLICATE = make_item(
    external_id="bad-dup",
    occurred_at="2026-01-20T12:00:00-06:00",
    amount="700.00",
    channel="BRANCH",
)


def _fixture(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


@pytest.fixture
def store(tmp_path: Path) -> LocalFsObjectStore:
    return LocalFsObjectStore(tmp_path / "objects")


def _items(data: bytes, content_type: str) -> list[TransactionInput]:
    rows = list(iter_rows(data, content_type, start=0, limit=100))
    items, errors = validate_batch(
        [row.data for row in rows if row.data is not None],
        now=clock.now(),
        max_future_minutes=10,
    )
    assert errors == []
    return items


def _xlsx(rows: Sequence[Sequence[object]]) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for row in rows:
        sheet.append(list(row))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _upload(
    tenant: CreatedTenant,
    store: LocalFsObjectStore,
    queue: PostgresJobQueue,
    data: bytes,
    content_type: str = CSV_TYPE,
) -> tuple[UUID, bool]:
    with staff_session(tenant.tenant_id, tenant.admin_user_id) as s:
        row, created = upload_file(
            s,
            store,
            data=data,
            content_type=content_type,
            uploaded_by="staff:test",
            queue=queue,
        )
        return row.id, created


def _job(tenant: CreatedTenant, file_id: UUID) -> Job:
    return Job(
        id=uuid7(),
        tenant_id=tenant.tenant_id,
        kind=INGEST_FILE_KIND,
        key=f"ingest_file:{file_id}",
        payload={"file_id": str(file_id)},
        attempts=1,
        max_attempts=3,
        lock_token=uuid7(),
    )


def _ingest_job(store: LocalFsObjectStore, queue: PostgresJobQueue) -> IngestFileJob:
    return IngestFileJob(
        session_factory=get_sessionmaker(),
        store=store,
        queue=queue,
        settings=get_settings(),
    )


def _run(
    tenant: CreatedTenant,
    store: LocalFsObjectStore,
    queue: PostgresJobQueue,
    file_id: UUID,
) -> None:
    with tenant_session(system_ctx(tenant.tenant_id)) as s:
        _ingest_job(store, queue).run(s, _job(tenant, file_id))


def _file(tenant: CreatedTenant, file_id: UUID) -> TransactionFile:
    with tenant_session(system_ctx(tenant.tenant_id)) as s:
        return s.get_one(TransactionFile, (tenant.tenant_id, file_id))


def _external_ids(tenant: CreatedTenant) -> list[str]:
    with tenant_session(system_ctx(tenant.tenant_id)) as s:
        return sorted(s.scalars(select(Transaction.external_id)).all())


def _report(store: LocalFsObjectStore, tenant: CreatedTenant, file_id: UUID) -> bytes:
    data = read_object(
        store, tenant.tenant_id, f"transactions/files/{file_id}.errors.csv"
    )
    assert data is not None
    return data


def _pre_ingest_the_duplicate(tenant: CreatedTenant, queue: PostgresJobQueue) -> None:
    with staff_session(tenant.tenant_id, tenant.admin_user_id) as s:
        report: BatchReport = ingest_items(s, [DUPLICATE], queue=queue)
    assert report.accepted == 1


BAD_ROWS_REPORT = (
    b"row,column,code\n4,amount,INVALID\n8,customer_ref,CUSTOMER_REF_UNKNOWN\n"
)
HEADER = (
    b"external_id,occurred_at,customer_ref,direction,rail,instrument,amount,currency\n"
)
GOOD_BAD_ROWS = [f"bad-{n:03d}" for n in (1, 2, 3, 5, 6, 9, 10)]


# ---- the readers ---------------------------------------------------------------------


def test_a_csv_a_jsonl_and_an_xlsx_of_the_same_rows_give_the_same_items():
    csv_items = _items(_fixture("deposits.csv"), CSV_TYPE)
    assert len(csv_items) == 4
    for name, content_type in (
        ("deposits_semicolon_bom.csv", CSV_TYPE),
        ("deposits.jsonl", JSONL_TYPE),
        ("deposits.xlsx", XLSX_TYPE),
    ):
        items = _items(_fixture(name), content_type)
        assert items == csv_items, name
        assert [content_hash(i) for i in items] == [content_hash(i) for i in csv_items]
    first = csv_items[0]
    assert first.counterparty is not None
    assert first.counterparty.name == "Pérez Gómez, Juan"
    assert first.raw == {"sucursal": "Sucursal Centro"}
    assert str(csv_items[1].amount) == "25000.50"


def test_bom_semicolons_and_quoted_commas_are_read():
    data = _fixture("deposits_semicolon_bom.csv")
    assert data.startswith(b"\xef\xbb\xbf")
    header = read_header(data, CSV_TYPE)
    assert header[0] == "external_id"
    assert "counterparty.name" in header
    rows = list(iter_rows(data, CSV_TYPE, start=0, limit=10))
    assert [row.number for row in rows] == [1, 2, 3, 4]
    names = [
        row.data["counterparty"]
        for row in rows
        if row.data and "counterparty" in row.data
    ]
    assert names == [
        {"kind": "PERSON", "name": "Pérez Gómez, Juan", "country": "MX"},
        {
            "kind": "ORGANIZATION",
            "name": "Comercial Ejemplo, SA de CV",
            "country": "MX",
        },
    ]
    # A semicolon inside quotes is a cell's text, not a delimiter.
    quoted = (
        b"external_id;occurred_at;customer_ref;direction;rail;instrument;amount;"
        b'currency;note\r\n"a;1";x;c;IN;CASH;CASH;1;MXN;"uno; dos"\r\n'
    )
    (row,) = iter_rows(quoted, CSV_TYPE, start=0, limit=5)
    assert row.data is not None
    assert row.data["external_id"] == "a;1"
    assert row.data["raw"] == {"note": "uno; dos"}


def test_dotted_headers_nest_and_unknown_columns_go_to_raw():
    item = row_to_item({
        "external_id": "tx-1",
        "counterparty.name": "Ana",
        "counterparty.identifier.id_type": "RFC",
        "counterparty.identifier.value": "XAXX010101000",
        "crypto.network": "ethereum",
        "device.ip": "203.0.113.7",
        "raw.canal_origen": "kiosko",
        "sucursal": "Centro",
        "fee_amount": "",
        "notas.extra": "x",
        "counterparty": "plano",
    })
    assert item == {
        "external_id": "tx-1",
        "counterparty": {
            "name": "Ana",
            "identifier": {"id_type": "RFC", "value": "XAXX010101000"},
        },
        "crypto": {"network": "ethereum"},
        "device": {"ip": "203.0.113.7"},
        # The blank fee is left out; a cell that would overwrite a nested field
        # goes to raw like any unknown column.
        "raw": {
            "canal_origen": "kiosko",
            "sucursal": "Centro",
            "notas.extra": "x",
            "counterparty": "plano",
        },
    }


def test_blank_rows_are_skipped_but_keep_their_number_and_pages_slice_rows():
    data = (
        HEADER
        + b"a,t,c,IN,CASH,CASH,1,MXN\n"
        + b",,,,,,,\n"
        + b"b,t,c,IN,CASH,CASH,1,MXN\n"
        + b"c,t,c,IN,CASH,CASH,1,MXN\n"
    )
    assert count_rows(data, CSV_TYPE) == 3
    page = list(iter_rows(data, CSV_TYPE, start=1, limit=5))
    assert [(row.number, row.data and row.data["external_id"]) for row in page] == [
        (3, "b"),
        (4, "c"),
    ]


def test_a_malformed_json_line_is_a_row_error_and_the_line_number_is_kept():
    data = b'{"external_id": "a"}\n\n[1, 2]\n{not json\n{"external_id": "b"}\n'
    assert read_header(data, JSONL_TYPE) == ["external_id"]
    rows = list(iter_rows(data, JSONL_TYPE, start=0, limit=10))
    assert [row.number for row in rows] == [1, 3, 4, 5]
    assert [row.data is None for row in rows] == [False, True, True, False]
    assert [(e.index, e.field, e.code.value) for e in rows[1].errors] == [
        (3, "", "INVALID")
    ]


def test_xlsx_cells_become_text():
    data = _xlsx([
        [
            "external_id",
            "occurred_at",
            "customer_ref",
            "direction",
            "rail",
            "instrument",
            "amount",
            "currency",
            "fee_amount",
        ],
        [
            "x-1",
            "2026-01-15T10:00:00-06:00",
            "cust-1",
            "IN",
            "CASH",
            "CASH",
            1500,
            "MXN",
            12.5,
        ],
        [None, None, None, None, None, None, None, None, None],
    ])
    (row,) = iter_rows(data, XLSX_TYPE, start=0, limit=5)
    assert row.data is not None
    assert (row.data["amount"], row.data["fee_amount"]) == ("1500", "12.5")


@pytest.mark.parametrize(
    ("data", "content_type"),
    [
        (b"", CSV_TYPE),
        (b"\n\n", CSV_TYPE),
        (b"\xff\xfe\x00bad", CSV_TYPE),
        (b"external_id,amount\nx,1\n", CSV_TYPE),
        (b"not a zip", XLSX_TYPE),
        (b"", JSONL_TYPE),
        (b"[1, 2]\n", JSONL_TYPE),
        (b"external_id\nx\n", "text/plain"),
    ],
)
def test_an_unreadable_or_headerless_file_is_invalid(data: bytes, content_type: str):
    with pytest.raises(FileInvalidError) as raised:
        _ = read_header(data, content_type)
    assert raised.value.code == "TRANSACTION_FILE_INVALID"


def test_an_unreadable_or_headerless_file_is_invalid_and_oversize_is_too_large(
    tenant: CreatedTenant,
    store: LocalFsObjectStore,
    queue: PostgresJobQueue,
    monkeypatch: pytest.MonkeyPatch,
):
    header_only = HEADER
    for data, content_type in (
        (b"external_id,amount\nx,1\n", CSV_TYPE),
        (header_only, CSV_TYPE),
        (_fixture("deposits.csv"), "application/pdf"),
        (_fixture("deposits.csv"), XLSX_TYPE),
    ):
        with pytest.raises(FileInvalidError):
            _ = _upload(tenant, store, queue, data, content_type)
    monkeypatch.setattr(get_settings(), "monitoring_file_max_rows", 3)
    with pytest.raises(FileTooLargeError) as raised:
        _ = _upload(tenant, store, queue, _fixture("deposits.csv"))
    assert (raised.value.status, raised.value.code) == (
        413,
        "TRANSACTION_FILE_TOO_LARGE",
    )
    monkeypatch.setattr(get_settings(), "monitoring_file_max_rows", 100)
    monkeypatch.setattr(get_settings(), "monitoring_file_max_mb", 1)
    big = header_only + b"x" * (1024 * 1024)
    with pytest.raises(FileTooLargeError):
        _ = _upload(tenant, store, queue, big)
    with tenant_session(system_ctx(tenant.tenant_id)) as s:
        assert s.scalar(select(func.count()).select_from(TransactionFile)) == 0


def test_the_error_report_is_row_column_and_code():
    report = error_report_csv([RowError(row=3, column="amount", code="INVALID")])
    assert report == b"row,column,code\n3,amount,INVALID\n"


# ---- the upload ----------------------------------------------------------------------


def test_an_upload_seals_the_file_stores_a_row_audits_and_enqueues_the_job(
    tenant: CreatedTenant, store: LocalFsObjectStore, queue: PostgresJobQueue
):
    data = _fixture("deposits.csv")
    file_id, created = _upload(tenant, store, queue, data, "text/csv; charset=utf-8")
    assert created
    row = _file(tenant, file_id)
    assert row.status is FileStatus.UPLOADED
    assert (row.rows_total, row.size, row.content_type) == (4, len(data), CSV_TYPE)
    assert row.storage_key == f"transactions/files/{file_id}"
    (ref,) = [
        r
        for r in store.list(tenant.tenant_id, row.storage_key)
        if r.key == row.storage_key
    ]
    assert ref.content_type == "application/octet-stream"
    sealed = store.get(ref)
    assert sealed != data
    assert b"cust-1" not in sealed
    with tenant_session(system_ctx(tenant.tenant_id)) as s:
        assert open_blob(s, tenant.tenant_id, sealed) == data
        jobs = s.scalars(select(JobRow).where(JobRow.kind == INGEST_FILE_KIND)).all()
        assert [(j.key, dict(j.payload)) for j in jobs] == [
            (f"ingest_file:{file_id}", {"file_id": str(file_id)})
        ]
    (audit,) = [
        e
        for e in audit_events(tenant.tenant_id)
        if e.action == "monitoring.file_uploaded"
    ]
    assert audit.entity_id == str(file_id)
    assert audit.detail == {
        "file_id": str(file_id),
        "size": len(data),
        "rows_total": 4,
        "content_type": CSV_TYPE,
    }


def test_the_same_file_twice_returns_the_existing_row(
    tenant: CreatedTenant, store: LocalFsObjectStore, queue: PostgresJobQueue
):
    data = _fixture("deposits.jsonl")
    first, created = _upload(tenant, store, queue, data, JSONL_TYPE)
    again, created_again = _upload(tenant, store, queue, data, JSONL_TYPE)
    assert (created, created_again) == (True, False)
    assert again == first
    with tenant_session(system_ctx(tenant.tenant_id)) as s:
        assert s.scalar(select(func.count()).select_from(TransactionFile)) == 1
        assert (
            s.scalar(
                select(func.count())
                .select_from(JobRow)
                .where(JobRow.kind == INGEST_FILE_KIND)
            )
            == 1
        )


# ---- the job -------------------------------------------------------------------------


def test_the_job_accepts_valid_rows_and_reports_the_rest(
    tenant: CreatedTenant,
    customer: CustomerRelationship,
    store: LocalFsObjectStore,
    queue: PostgresJobQueue,
):
    _pre_ingest_the_duplicate(tenant, queue)
    file_id, _ = _upload(tenant, store, queue, _fixture("bad_rows.csv"))
    # Through the worker: the queued job, claimed and run by the runner.
    register_job_handler(INGEST_FILE_KIND, _ingest_job(store, queue).run)
    runner = JobRunner(
        queue, get_sessionmaker(), worker_id="w-files", kinds=[INGEST_FILE_KIND]
    )
    assert runner.run_once() == 1
    row = _file(tenant, file_id)
    assert row.status is FileStatus.COMPLETED
    assert row.finished_at is not None
    assert (
        row.rows_total,
        row.rows_accepted,
        row.rows_duplicate,
        row.rows_rejected,
    ) == (
        10,
        7,
        1,
        2,
    )
    assert row.cursor == 10
    assert row.errors_key == f"transactions/files/{file_id}.errors.csv"
    assert _report(store, tenant, file_id) == BAD_ROWS_REPORT
    assert _external_ids(tenant) == sorted([*GOOD_BAD_ROWS, "bad-dup"])
    with tenant_session(system_ctx(tenant.tenant_id)) as s:
        batches = s.scalars(
            select(TransactionBatch).where(TransactionBatch.file_id == file_id)
        ).all()
        job = s.scalars(select(JobRow).where(JobRow.kind == INGEST_FILE_KIND)).one()
    assert [(b.source, b.submitted_by, b.accepted, b.duplicates) for b in batches] == [
        (BatchSource.FILE, "staff:test", 7, 1)
    ]
    assert job.status.value == "DONE"


def test_row_errors_never_echo_cell_values(
    tenant: CreatedTenant,
    customer: CustomerRelationship,
    store: LocalFsObjectStore,
    queue: PostgresJobQueue,
    caplog: pytest.LogCaptureFixture,
):
    data = _fixture("bad_rows.csv")
    reader = csv.DictReader(io.StringIO(data.decode()))
    cells = {
        row[column]
        for row in reader
        for column in (
            "external_id",
            "customer_ref",
            "amount",
            "occurred_at",
            "counterparty.name",
        )
        if row[column]
    }
    assert {"doce pesos", "cust-fantasma", "Martínez Ruiz, Ana"} <= cells
    file_id, _ = _upload(tenant, store, queue, data)
    caplog.set_level(logging.DEBUG)
    _run(tenant, store, queue, file_id)
    texts = [
        _report(store, tenant, file_id).decode(),
        _file(tenant, file_id).last_error or "",
        caplog.text,
        "\n".join(str(r.__dict__) for r in caplog.records),
    ]
    assert [
        (cell, i) for cell in cells for i, text in enumerate(texts) if cell in text
    ] == []


def test_the_job_resumes_from_the_cursor_after_a_crash(
    tenant: CreatedTenant,
    customer: CustomerRelationship,
    store: LocalFsObjectStore,
    queue: PostgresJobQueue,
    monkeypatch: pytest.MonkeyPatch,
):
    _pre_ingest_the_duplicate(tenant, queue)
    monkeypatch.setattr(get_settings(), "monitoring_file_chunk_rows", 3)
    file_id, _ = _upload(tenant, store, queue, _fixture("bad_rows.csv"))
    real_page = files.ingest_page
    pages: list[int] = []

    def crash_on_the_third_page(
        session: Session,
        file: TransactionFile,
        rows: Sequence[files.FileRow],
        *,
        store: ObjectStore,
        queue: PostgresJobQueue,
    ) -> tuple[int, int, list[RowError]]:
        pages.append(rows[0].number)
        if len(pages) == 3:
            msg = "the worker died"
            raise RuntimeError(msg)
        return real_page(session, file, rows, store=store, queue=queue)

    monkeypatch.setattr(files, "ingest_page", crash_on_the_third_page)
    with pytest.raises(RuntimeError):
        _run(tenant, store, queue, file_id)
    crashed = _file(tenant, file_id)
    assert crashed.status is FileStatus.FAILED
    assert crashed.last_error == "INTERNAL: RuntimeError"
    assert crashed.cursor == 6
    assert (crashed.rows_accepted, crashed.rows_rejected) == (5, 1)
    assert _report(store, tenant, file_id) == b"row,column,code\n4,amount,INVALID\n"

    monkeypatch.setattr(files, "ingest_page", real_page)
    _run(tenant, store, queue, file_id)
    assert pages == [1, 4, 7]
    done = _file(tenant, file_id)
    assert done.status is FileStatus.COMPLETED
    assert done.last_error is None
    assert (
        done.cursor,
        done.rows_accepted,
        done.rows_duplicate,
        done.rows_rejected,
    ) == (
        10,
        7,
        1,
        2,
    )
    assert _external_ids(tenant) == sorted([*GOOD_BAD_ROWS, "bad-dup"])
    assert _report(store, tenant, file_id) == BAD_ROWS_REPORT
    # A completed file is not walked again.
    _run(tenant, store, queue, file_id)
    assert _file(tenant, file_id).rows_accepted == 7


def test_completion_publishes_the_file_event(
    tenant: CreatedTenant,
    customer: CustomerRelationship,
    store: LocalFsObjectStore,
    queue: PostgresJobQueue,
):
    file_id, _ = _upload(tenant, store, queue, _fixture("deposits.xlsx"), XLSX_TYPE)
    _run(tenant, store, queue, file_id)
    with tenant_session(system_ctx(tenant.tenant_id)) as s:
        (row,) = s.scalars(
            select(OutboxRow).where(OutboxRow.type == TRANSACTION_FILE_COMPLETED)
        ).all()
        event = event_from_row(row)
    assert event.subject == f"transaction_file/{file_id}"
    assert event.key == str(file_id)
    assert event.data == {
        "file_id": str(file_id),
        "status": "COMPLETED",
        "rows_accepted": 4,
        "rows_rejected": 0,
    }
    assert _file(tenant, file_id).errors_key is None


def test_a_chunk_is_one_transaction(
    tenant: CreatedTenant,
    customer: CustomerRelationship,
    store: LocalFsObjectStore,
    queue: PostgresJobQueue,
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.setattr(get_settings(), "monitoring_file_chunk_rows", 3)
    file_id, _ = _upload(tenant, store, queue, _fixture("bad_rows.csv"))
    real_batch = ingest_batch

    def fail_after_writing_bad_009(
        session: Session,
        items: Sequence[TransactionInput],
        *,
        source: BatchSource,
        submitted_by: str,
        queue: PostgresJobQueue,
        file_id: UUID | None = None,
    ) -> BatchReport:
        report = real_batch(
            session,
            items,
            source=source,
            submitted_by=submitted_by,
            queue=queue,
            file_id=file_id,
        )
        if any(item.external_id == "bad-009" for item in items):
            session.flush()
            msg = "the database went away"
            raise RuntimeError(msg)
        return report

    monkeypatch.setattr(files, "ingest_batch", fail_after_writing_bad_009)
    with pytest.raises(RuntimeError):
        _run(tenant, store, queue, file_id)
    # Pages 0 and 1 (rows 1 to 6) committed; page 2 (rows 7 to 9) left nothing.
    assert _external_ids(tenant) == [
        "bad-001",
        "bad-002",
        "bad-003",
        "bad-005",
        "bad-006",
    ]
    failed = _file(tenant, file_id)
    assert (failed.cursor, failed.status) == (6, FileStatus.FAILED)
    with tenant_session(system_ctx(tenant.tenant_id)) as s:
        assert (
            s.scalar(
                select(func.count())
                .select_from(TransactionBatch)
                .where(TransactionBatch.file_id == file_id)
            )
            == 2
        )


def test_a_page_whose_rows_are_all_refused_ingests_nothing(
    tenant: CreatedTenant,
    customer: CustomerRelationship,
    store: LocalFsObjectStore,
    queue: PostgresJobQueue,
):
    lines = [json.dumps(make_item(external_id="j-1", customer_ref="nadie")), "{roto"]
    data = "\n".join(lines).encode()
    file_id, _ = _upload(tenant, store, queue, data, JSONL_TYPE)
    _run(tenant, store, queue, file_id)
    row = _file(tenant, file_id)
    assert (row.status, row.rows_accepted, row.rows_rejected) == (
        FileStatus.COMPLETED,
        0,
        2,
    )
    assert _report(store, tenant, file_id) == (
        b"row,column,code\n1,customer_ref,CUSTOMER_REF_UNKNOWN\n2,,INVALID\n"
    )
    assert _external_ids(tenant) == []


def test_a_known_external_id_with_other_content_is_reported_and_the_rest_ingested(
    tenant: CreatedTenant,
    customer: CustomerRelationship,
    store: LocalFsObjectStore,
    queue: PostgresJobQueue,
):
    changed = {**DUPLICATE, "amount": "701.00"}
    with staff_session(tenant.tenant_id, tenant.admin_user_id) as s:
        _ = ingest_items(s, [changed], queue=queue)
    file_id, _ = _upload(tenant, store, queue, _fixture("bad_rows.csv"))
    _run(tenant, store, queue, file_id)
    row = _file(tenant, file_id)
    assert (row.rows_accepted, row.rows_duplicate, row.rows_rejected) == (7, 0, 3)
    assert _report(store, tenant, file_id) == (
        b"row,column,code\n4,amount,INVALID\n7,external_id,EXTERNAL_ID_CONFLICT\n"
        b"8,customer_ref,CUSTOMER_REF_UNKNOWN\n"
    )


def test_a_repeated_external_id_in_a_page_is_a_row_error(
    tenant: CreatedTenant,
    customer: CustomerRelationship,
    store: LocalFsObjectStore,
    queue: PostgresJobQueue,
):
    item = json.dumps({**make_item(external_id="r-1"), "channel": "APP"})
    file_id, _ = _upload(tenant, store, queue, f"{item}\n{item}\n".encode(), JSONL_TYPE)
    _run(tenant, store, queue, file_id)
    assert (
        _report(store, tenant, file_id)
        == b"row,column,code\n2,external_id,DUPLICATE_IN_BATCH\n"
    )
    assert _file(tenant, file_id).rows_accepted == 1


def test_a_missing_upload_fails_the_file(
    tenant: CreatedTenant,
    customer: CustomerRelationship,
    store: LocalFsObjectStore,
    queue: PostgresJobQueue,
):
    file_id, _ = _upload(tenant, store, queue, _fixture("deposits.csv"))
    row = _file(tenant, file_id)
    (ref,) = [
        r
        for r in store.list(tenant.tenant_id, row.storage_key)
        if r.key == row.storage_key
    ]
    store.delete(ref)
    with pytest.raises(NotFoundError):
        _run(tenant, store, queue, file_id)
    failed = _file(tenant, file_id)
    assert (failed.status, failed.last_error) == (
        FileStatus.FAILED,
        "NOT_FOUND: NotFoundError",
    )


def test_another_run_that_moved_the_cursor_stops_this_one(
    tenant: CreatedTenant,
    customer: CustomerRelationship,
    store: LocalFsObjectStore,
    queue: PostgresJobQueue,
    monkeypatch: pytest.MonkeyPatch,
):
    file_id, _ = _upload(tenant, store, queue, _fixture("deposits.csv"))
    real_rows = files.iter_rows

    def advance_then_read(
        data: bytes, content_type: str, *, start: int, limit: int
    ) -> Generator[files.FileRow]:
        # Another run commits a page between this run's open and its first page.
        with tenant_session(system_ctx(tenant.tenant_id)) as s:
            s.get_one(TransactionFile, (tenant.tenant_id, file_id)).cursor = 2
        return real_rows(data, content_type, start=start, limit=limit)

    monkeypatch.setattr(files, "iter_rows", advance_then_read)
    with pytest.raises(files.FileCursorMovedError):
        _run(tenant, store, queue, file_id)
    assert _external_ids(tenant) == []
