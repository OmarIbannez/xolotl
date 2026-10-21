"""File ingest: CSV, JSON Lines and XLSX uploads walked in pages (SP-04 spec §6.2).

``upload_file`` checks a file's size, type, header and row count, seals it under
the tenant DATA key at ``transactions/files/<id>``, writes the ``UPLOADED`` row,
audits the upload and enqueues ``monitoring.ingest_file``. ``IngestFileJob`` walks
the file ``MONITORING_FILE_CHUNK_ROWS`` rows at a time from ``cursor``, each page
in a transaction of its own (``ingest_page``): the rows that validate are ingested
as one ``FILE`` batch, the others are reported and counted, so one bad row never
bounces the file (Decision 5). A crash leaves every committed page in place and the
retry resumes from ``cursor``.

The row error report ``transactions/files/<id>.errors.csv`` holds ``row``,
``column`` and ``code`` only, never a cell value (plan 4a Deviation 12: each page
rewrites it, keeping the earlier pages' rows). It sits beside the upload rather
than under it (the plan's ``<id>/errors.csv``): the upload's key is an object, so
it cannot also be a folder of the object store. Nothing here logs a cell, and every
error raised over the file's content carries a fixed sentence.

Reading: a CSV is UTF-8 (a BOM is tolerated) with a comma or semicolon delimiter
chosen from the header line; an XLSX is read from its first sheet; a JSON Lines
file is one object per line. Blank rows and lines are skipped but keep their place
in the numbering, so a row number points at the row a person sees.
"""

import csv
import hashlib
import io
import json
import logging
import re
import zipfile
import zlib
from collections.abc import Callable, Generator, Iterator, Mapping, Sequence
from contextlib import closing
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from itertools import islice
from typing import ClassVar, cast
from uuid import UUID

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.orm import Session

from xolotl.modules.monitoring.enums import BatchSource, FileStatus
from xolotl.modules.monitoring.errors import (
    BatchInvalidError,
    ExternalIdConflictError,
    FileInvalidError,
    FileTooLargeError,
)
from xolotl.modules.monitoring.events import (
    TRANSACTION_FILE_COMPLETED,
    publish_monitoring_event,
)
from xolotl.modules.monitoring.kinds import INGEST_FILE_KIND
from xolotl.modules.monitoring.txn.ingest import ingest_batch
from xolotl.modules.monitoring.txn.models import TransactionFile
from xolotl.modules.monitoring.txn.schema import (
    FieldError,
    FieldErrorCode,
    TransactionInput,
    validate_batch,
)
from xolotl.platform import clock
from xolotl.platform.audit import AuditEvent, record
from xolotl.platform.config import Settings, get_settings
from xolotl.platform.crypto.blobs import open_blob, seal_blob
from xolotl.platform.db import SessionFactory, bound_context, tenant_session
from xolotl.platform.errors import NotFoundError, XolotlError
from xolotl.platform.ids import uuid7
from xolotl.platform.jobs.queue import Job, PostgresJobQueue
from xolotl.platform.storage import ObjectStore
from xolotl.platform.tenancy.context import TenantContext

CSV_TYPE = "text/csv"
JSONL_TYPE = "application/x-ndjson"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ACCEPTED_TYPES: frozenset[str] = frozenset({CSV_TYPE, JSONL_TYPE, XLSX_TYPE})
REQUIRED_COLUMNS: tuple[str, ...] = (
    "external_id",
    "occurred_at",
    "customer_ref",
    "direction",
    "rail",
    "instrument",
    "amount",
    "currency",
)
SEALED_CONTENT_TYPE = "application/octet-stream"
ERRORS_HEADER: tuple[str, str, str] = ("row", "column", "code")
# The widest XLSX row read: every field and nested field, with room for extras.
MAX_XLSX_COLUMNS = 256
# How many times ``monitoring_file_max_mb`` an XLSX may unpack to: spreadsheets
# usually compress 5 to 10 times, and the row limit bounds an honest file anyway.
XLSX_MAX_EXPANSION = 10
CONFLICT_CODE = "EXTERNAL_ID_CONFLICT"
_MIB = 1024 * 1024
_LAST_ERROR_LIMIT = 200
_FIELDS: frozenset[str] = frozenset(TransactionInput.model_fields)
# Where a refused batch item's index and field sit in its error's ``loc``.
_LOC_INDEX = 1
_LOC_FIELD = 2
_CONFLICT_DETAIL = re.compile(r"transactions\[(\d+)\]")
_XLSX_ERRORS = (
    InvalidFileException,
    zipfile.BadZipFile,
    zlib.error,
    # The XML parsers' ParseError (ElementTree's and lxml's) is a SyntaxError.
    SyntaxError,
    EOFError,
    KeyError,
    ValueError,
    OSError,
)
_UNREADABLE = "the file cannot be read"
_NO_HEADER = "the file has no header row"
_MISSING_COLUMN = "a required column is missing"
_NOT_AN_OBJECT = "the first line is not a JSON object"
_NO_ROWS = "the file has no data rows"
_WRONG_TYPE = "the content type is not CSV, JSON Lines or XLSX"

logger = logging.getLogger(__name__)


@dataclass(frozen=True, slots=True)
class FileRow:
    """One data row: ``number`` is the 1-based data row or the JSON Lines line.

    ``data`` is the item the row spells, or ``None`` with ``errors`` saying why
    when the row cannot be read at all (a malformed JSON line).
    """

    number: int
    data: dict[str, object] | None
    errors: tuple[FieldError, ...]


@dataclass(frozen=True, slots=True)
class RowError:
    """One rejected field of a file row: where and why, never the value."""

    row: int
    column: str
    code: str


# ---- reading -------------------------------------------------------------------------


def media_type(content_type: str) -> str:
    """Drop a content type's parameters (``; charset=utf-8``) and its case.

    Returns:
        The bare media type.

    """
    return content_type.partition(";")[0].strip().lower()


def _text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise FileInvalidError(_UNREADABLE) from None


def _csv_records(data: bytes) -> Generator[list[str]]:
    text = _text(data)
    first_line = text.split("\n", 1)[0]
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    try:
        yield from csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    except csv.Error:
        raise FileInvalidError(_UNREADABLE) from None


def _xlsx_cell(value: object) -> str:
    # Dates as ISO text, whole numbers without a decimal point, the rest as text.
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _xlsx_records(data: bytes) -> Generator[list[str]]:
    # A read-only workbook parses its sheet lazily, so a corrupt sheet fails while
    # rows are read; it still loads the shared strings whole before the first row,
    # so a file whose parts unpack past the ceiling is refused unparsed.
    ceiling = get_settings().monitoring_file_max_mb * _MIB * XLSX_MAX_EXPANSION
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            unpacked = sum(info.file_size for info in archive.infolist())
    except _XLSX_ERRORS:
        raise FileInvalidError(_UNREADABLE) from None
    if unpacked > ceiling:
        raise FileInvalidError(_UNREADABLE)
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except _XLSX_ERRORS:
        raise FileInvalidError(_UNREADABLE) from None
    try:
        sheets = book.worksheets
        if not sheets:
            raise FileInvalidError(_NO_HEADER)
        for row in sheets[0].iter_rows(values_only=True, max_col=MAX_XLSX_COLUMNS):
            yield [_xlsx_cell(value) for value in row]
    except _XLSX_ERRORS:
        raise FileInvalidError(_UNREADABLE) from None
    finally:
        book.close()


def _header_of(records: Iterator[list[str]]) -> list[str]:
    # The header is the first non-blank record.
    for line in records:
        cells = [cell.strip() for cell in line]
        if any(cells):
            return _header_names(cells)
    raise FileInvalidError(_NO_HEADER)


def _table_rows(
    data: bytes, content_type: str
) -> Generator[tuple[int, dict[str, str]]]:
    # Each non-blank record after the header, numbered by its place after it.
    records = _xlsx_records(data) if content_type == XLSX_TYPE else _csv_records(data)
    with closing(records):
        names = _header_of(records)
        for number, line in enumerate(records, start=1):
            cells = [cell.strip() for cell in line]
            if any(cells):
                yield (
                    number,
                    {
                        name: cell
                        for name, cell in zip(names, cells, strict=False)
                        if name and cell
                    },
                )


def _header_names(cells: list[str]) -> list[str]:
    if any(column not in cells for column in REQUIRED_COLUMNS):
        raise FileInvalidError(_MISSING_COLUMN)
    return cells


def _jsonl_lines(data: bytes) -> Generator[tuple[int, str]]:
    # Split on line feeds only: a JSON string may hold U+2028 and other breaks.
    for number, line in enumerate(_text(data).split("\n"), start=1):
        if line.strip():
            yield number, line


def _json_object(line: str) -> dict[str, object] | None:
    try:
        value = cast("object", json.loads(line, parse_float=Decimal))
    except ValueError:
        return None
    if not isinstance(value, dict):
        return None
    # JSON object keys are always text.
    return cast("dict[str, object]", value)


def _all_rows(data: bytes, content_type: str) -> Generator[FileRow]:
    if content_type == JSONL_TYPE:
        for number, line in _jsonl_lines(data):
            item = _json_object(line)
            if item is None:
                error = FieldError(index=number, field="", code=FieldErrorCode.INVALID)
                yield FileRow(number=number, data=None, errors=(error,))
            else:
                yield FileRow(number=number, data=item, errors=())
        return
    rows = _table_rows(data, content_type)
    with closing(rows):
        for number, cells in rows:
            yield FileRow(number=number, data=row_to_item(cells), errors=())


def _require_type(content_type: str) -> str:
    if content_type not in ACCEPTED_TYPES:
        raise FileInvalidError(_WRONG_TYPE)
    return content_type


def read_header(data: bytes, content_type: str) -> list[str]:
    """Read a file's column names (for JSON Lines, the first object's keys).

    Returns:
        The header cells, trimmed, in file order.

    Raises:
        FileInvalidError: When the type is not accepted, the file cannot be read,
            has no header, lacks one of ``REQUIRED_COLUMNS`` (CSV and XLSX), or its
            first line is not a JSON object (JSON Lines).

    """
    content_type = _require_type(content_type)
    if content_type == JSONL_TYPE:
        lines = _jsonl_lines(data)
        with closing(lines):
            first = next(lines, None)
        if first is None:
            raise FileInvalidError(_NO_HEADER)
        item = _json_object(first[1])
        if item is None:
            raise FileInvalidError(_NOT_AN_OBJECT)
        return list(item)
    records = _xlsx_records(data) if content_type == XLSX_TYPE else _csv_records(data)
    with closing(records):
        return _header_of(records)


def count_rows(data: bytes, content_type: str) -> int:
    """Count a file's data rows (non-blank rows after the header, or lines).

    Returns:
        How many rows ``iter_rows`` yields over the whole file.

    """
    return sum(1 for _ in _all_rows(data, _require_type(content_type)))


def iter_rows(
    data: bytes, content_type: str, *, start: int, limit: int
) -> Generator[FileRow]:
    """Yield the data rows ``[start, start + limit)`` of a file, counted from 0.

    Yields:
        Each row in file order, as an item or the errors that kept it from one.

    """
    rows = _all_rows(data, _require_type(content_type))
    with closing(rows):
        yield from islice(rows, start, start + limit)


def row_to_item(cells: Mapping[str, str]) -> dict[str, object]:
    """Turn one CSV or XLSX row into the item it spells.

    A dotted header nests (``counterparty.name``, ``crypto.network``); a blank cell
    is left out; ``raw.<key>`` and any column that is not a field of the canonical
    transaction go into ``raw`` under the key or the column name (plan 4a
    Deviation 12). A cell that would overwrite another one goes into ``raw`` too.

    Returns:
        The item, as ``parse_item`` takes it.

    """
    item: dict[str, object] = {}
    raw: dict[str, object] = {}
    nodes: dict[tuple[str, ...], dict[str, object]] = {(): item}

    def node(path: tuple[str, ...]) -> dict[str, object] | None:
        if path in nodes:
            return nodes[path]
        parent = node(path[:-1])
        if parent is None or path[-1] in parent:
            return None
        child: dict[str, object] = {}
        parent[path[-1]] = child
        nodes[path] = child
        return child

    for column, cell in cells.items():
        value = cell.strip()
        if not value:
            continue
        head, _, rest = column.partition(".")
        if head == "raw" and rest:
            raw[rest] = value
            continue
        if head not in _FIELDS or head == "raw":
            raw[column] = value
            continue
        path = (head, *rest.split(".")) if rest else (head,)
        parent = node(path[:-1])
        if parent is None or path[-1] in parent:
            raw[column] = value
        else:
            parent[path[-1]] = value
    if raw:
        item["raw"] = raw
    return item


# ---- the error report ----------------------------------------------------------------


def error_report_csv(errors: Sequence[RowError]) -> bytes:
    """Write row errors as CSV with the header ``row,column,code``.

    Returns:
        The report's UTF-8 bytes.

    """
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow(ERRORS_HEADER)
    writer.writerows((error.row, error.column, error.code) for error in errors)
    return out.getvalue().encode()


def _read_report(data: bytes) -> list[RowError]:
    rows = list(csv.reader(io.StringIO(data.decode(), newline="")))
    return [
        RowError(row=int(row[0]), column=row[1], code=row[2])
        for row in rows[1:]
        if len(row) == len(ERRORS_HEADER)
    ]


def errors_key_for(file_id: UUID) -> str:
    """Name the object that holds a file's row error report.

    Returns:
        ``transactions/files/<file_id>.errors.csv``.

    """
    return f"{storage_key_for(file_id)}.errors.csv"


def storage_key_for(file_id: UUID) -> str:
    """Name the object that holds a file's sealed upload.

    Returns:
        ``transactions/files/<file_id>``.

    """
    return f"transactions/files/{file_id}"


def read_object(store: ObjectStore, tenant_id: UUID, key: str) -> bytes | None:
    """Read the object stored under exactly ``key``.

    Returns:
        Its bytes, or ``None`` when there is none.

    """
    for ref in store.list(tenant_id, key):
        if ref.key == key:
            return store.get(ref)
    return None


# ---- the upload ----------------------------------------------------------------------


def upload_file(
    session: Session,
    store: ObjectStore,
    *,
    data: bytes,
    content_type: str,
    uploaded_by: str,
    queue: PostgresJobQueue,
) -> tuple[TransactionFile, bool]:
    """Check, seal and store an upload, and enqueue its ingest job.

    A file whose sha256 the tenant already uploaded is not stored again.

    Returns:
        The file's row and whether this call created it.

    Raises:
        FileTooLargeError: Over ``MONITORING_FILE_MAX_MB`` MiB or
            ``MONITORING_FILE_MAX_ROWS`` data rows.
        FileInvalidError: For a type outside ``ACCEPTED_TYPES``, an unreadable or
            headerless file, a missing column, or no data row.

    """
    settings = get_settings()
    if len(data) > settings.monitoring_file_max_mb * _MIB:
        raise FileTooLargeError
    content_type = _require_type(media_type(content_type))
    _ = read_header(data, content_type)
    rows_total = count_rows(data, content_type)
    if rows_total > settings.monitoring_file_max_rows:
        raise FileTooLargeError
    if rows_total == 0:
        raise FileInvalidError(_NO_ROWS)
    digest = hashlib.sha256(data).digest()
    existing = session.scalar(
        select(TransactionFile).where(TransactionFile.sha256 == digest)
    )
    if existing is not None:
        return existing, False
    tenant_id = bound_context(session).require_tenant_id()
    file_id = uuid7()
    key = storage_key_for(file_id)
    _ = store.put(
        tenant_id,
        key,
        seal_blob(session, tenant_id, data),
        content_type=SEALED_CONTENT_TYPE,
        metadata={"content_type": content_type},
    )
    row = TransactionFile(
        tenant_id=tenant_id,
        id=file_id,
        storage_key=key,
        sha256=digest,
        size=len(data),
        content_type=content_type,
        status=FileStatus.UPLOADED,
        rows_total=rows_total,
        uploaded_by=uploaded_by,
    )
    session.add(row)
    session.flush()
    _ = queue.enqueue(
        session,
        kind=INGEST_FILE_KIND,
        payload={"file_id": str(file_id)},
        key=f"ingest_file:{file_id}",
    )
    _ = record(
        session,
        AuditEvent(
            action="monitoring.file_uploaded",
            entity_type="transaction_file",
            entity_id=str(file_id),
            detail={
                "file_id": str(file_id),
                "size": len(data),
                "rows_total": rows_total,
                "content_type": content_type,
            },
        ),
    )
    return row, True


# ---- one page ------------------------------------------------------------------------


def _refused(exc: BatchInvalidError) -> dict[int, list[tuple[str, str]]]:
    # Each item is located ("transactions", index, field) with its code as type.
    found: dict[int, list[tuple[str, str]]] = {}
    for error in exc.errors:
        loc = error.loc
        index = loc[_LOC_INDEX] if len(loc) > _LOC_INDEX else None
        if not isinstance(index, int):
            raise exc
        field = str(loc[_LOC_FIELD]) if len(loc) > _LOC_FIELD else ""
        found.setdefault(index, []).append((field, error.type))
    if not found:
        raise exc
    return found


def _conflicting(exc: ExternalIdConflictError) -> int:
    match = _CONFLICT_DETAIL.fullmatch(exc.detail)
    if match is None:
        raise exc
    return int(match.group(1))


def _ingest_valid(
    session: Session,
    file: TransactionFile,
    pending: list[tuple[int, TransactionInput]],
    *,
    queue: PostgresJobQueue,
    errors: list[RowError],
) -> tuple[int, int]:
    # A refused batch names the failing items: report them, drop them and ingest
    # the rest again, each attempt under a savepoint so a refusal leaves nothing.
    while pending:
        try:
            with session.begin_nested():
                report = ingest_batch(
                    session,
                    [item for _, item in pending],
                    source=BatchSource.FILE,
                    submitted_by=file.uploaded_by,
                    queue=queue,
                    file_id=file.id,
                )
        except BatchInvalidError as exc:
            refused = _refused(exc)
            errors.extend(
                RowError(row=pending[index][0], column=field, code=code)
                for index, fields in refused.items()
                for field, code in fields
            )
            pending = [p for index, p in enumerate(pending) if index not in refused]
        except ExternalIdConflictError as exc:
            index = _conflicting(exc)
            errors.append(
                RowError(
                    row=pending[index][0], column="external_id", code=CONFLICT_CODE
                )
            )
            del pending[index]
        else:
            return report.accepted, report.duplicates
    return 0, 0


def _rewrite_report(
    store: ObjectStore, file: TransactionFile, first_row: int, errors: list[RowError]
) -> None:
    # The rows of earlier pages are kept; any row of this page or later is from an
    # attempt that did not commit, and goes.
    key = errors_key_for(file.id)
    previous = read_object(store, file.tenant_id, key)
    if previous is None and not errors:
        return
    kept = [] if previous is None else _read_report(previous)
    report = [error for error in kept if error.row < first_row] + errors
    _ = store.put(file.tenant_id, key, error_report_csv(report), content_type=CSV_TYPE)
    file.errors_key = key if report else None


def ingest_page(
    session: Session,
    file: TransactionFile,
    rows: Sequence[FileRow],
    *,
    store: ObjectStore,
    queue: PostgresJobQueue,
) -> tuple[int, int, list[RowError]]:
    """Ingest one page of a file in the caller's transaction and advance ``file``.

    Rows that validate go through ``ingest_batch`` as one ``FILE`` batch; a row
    that fails validation, repeats an ``external_id`` of the page, names an
    unknown customer or another customer's account, or re-sends a known
    ``external_id`` with other content is reported instead. The counters and
    ``cursor`` of ``file`` advance and the error report is rewritten.

    Returns:
        The accepted and duplicate counts and the page's row errors.

    """
    errors = [
        RowError(row=row.number, column=error.field, code=error.code.value)
        for row in rows
        for error in row.errors
    ]
    numbers = [row.number for row in rows if row.data is not None]
    items, field_errors = validate_batch(
        [row.data for row in rows if row.data is not None],
        now=clock.now(),
        max_future_minutes=get_settings().monitoring_max_future_minutes,
    )
    failed = {error.index for error in field_errors}
    errors.extend(
        RowError(row=numbers[error.index], column=error.field, code=error.code.value)
        for error in field_errors
    )
    valid = [index for index in range(len(numbers)) if index not in failed]
    pending = [(numbers[index], item) for index, item in zip(valid, items, strict=True)]
    accepted, duplicates = _ingest_valid(
        session, file, pending, queue=queue, errors=errors
    )
    errors.sort(key=lambda error: (error.row, error.column))
    file.rows_accepted += accepted
    file.rows_duplicate += duplicates
    file.rows_rejected += len({error.row for error in errors})
    file.cursor += len(rows)
    if rows:
        _rewrite_report(store, file, rows[0].number, errors)
    return accepted, duplicates, errors


# ---- the job -------------------------------------------------------------------------


class FileCursorMovedError(XolotlError):
    """Another run advanced the file while this one held its rows."""

    default_code: ClassVar[str] = "TRANSACTION_FILE_CURSOR_MOVED"
    default_detail: ClassVar[str] = "another run advanced the file"


def _last_error(exc: BaseException) -> str:
    # The code and the class only: an exception's text may quote what it read.
    code = exc.code if isinstance(exc, XolotlError) else "INTERNAL"
    return f"{code}: {type(exc).__name__}"[:_LAST_ERROR_LIMIT]


@dataclass(frozen=True, slots=True)
class _Opened:
    data: bytes
    content_type: str
    cursor: int
    rows_total: int


class IngestFileJob:
    """The ``monitoring.ingest_file`` handler: a file's pages from ``cursor`` on."""

    def __init__(
        self,
        *,
        session_factory: SessionFactory,
        store: ObjectStore,
        queue: PostgresJobQueue,
        settings: Settings,
        clock_fn: Callable[[], datetime] = clock.now,
    ) -> None:
        """Keep what each page needs."""
        self._session_factory: SessionFactory = session_factory
        self._store: ObjectStore = store
        self._queue: PostgresJobQueue = queue
        self._settings: Settings = settings
        self._clock: Callable[[], datetime] = clock_fn

    def run(self, session: Session, job: Job) -> None:
        """Ingest the job's file page by page, then mark it ``COMPLETED``.

        The runner's ``session`` only takes the final status and the event; every
        page is a transaction of its own. A failure marks the file ``FAILED`` in a
        fresh transaction and is raised again, so the queue retries from
        ``cursor``.
        """
        file_id = UUID(str(job.payload["file_id"]))
        ctx = bound_context(session)
        try:
            opened = self._open(ctx, file_id)
            if opened is None:
                return
            self._pages(ctx, file_id, opened, job)
        except Exception as exc:
            with tenant_session(ctx, factory=self._session_factory) as s:
                file = self._lock(s, file_id)
                file.status = FileStatus.FAILED
                file.last_error = _last_error(exc)
            raise
        # Outside the handler above: the runner's session holds the file's lock
        # here, so marking it FAILED from another session would wait on itself. A
        # failure now rolls back with the runner's session; the retry finds the
        # cursor at the end and completes.
        self._complete(session, file_id)

    @staticmethod
    def _lock(session: Session, file_id: UUID) -> TransactionFile:
        tenant_id = bound_context(session).require_tenant_id()
        return session.get_one(
            TransactionFile,
            (tenant_id, file_id),
            with_for_update=True,
            populate_existing=True,
        )

    def _open(self, ctx: TenantContext, file_id: UUID) -> _Opened | None:
        with tenant_session(ctx, factory=self._session_factory) as s:
            file = self._lock(s, file_id)
            if file.status is FileStatus.COMPLETED:
                return None
            file.status = FileStatus.PROCESSING
            file.last_error = None
            sealed = read_object(self._store, file.tenant_id, file.storage_key)
            if sealed is None:
                msg = "the uploaded file is missing"
                raise NotFoundError(msg)
            opened = _Opened(
                data=open_blob(s, file.tenant_id, sealed),
                content_type=file.content_type,
                cursor=file.cursor,
                rows_total=file.rows_total or 0,
            )
        logger.info(
            "transaction file ingest started",
            extra={
                "file_id": str(file_id),
                "cursor": opened.cursor,
                "rows_total": opened.rows_total,
            },
        )
        return opened

    def _pages(
        self, ctx: TenantContext, file_id: UUID, opened: _Opened, job: Job
    ) -> None:
        # One reader for the whole run, so a page does not re-read the rows before
        # it; each page checks under its lock that no other run moved the cursor.
        chunk = self._settings.monitoring_file_chunk_rows
        lease = timedelta(seconds=self._settings.monitoring_process_lease_seconds)
        cursor = opened.cursor
        rows = iter_rows(
            opened.data,
            opened.content_type,
            start=cursor,
            limit=max(opened.rows_total - cursor, 0),
        )
        with closing(rows):
            while cursor < opened.rows_total:
                page = list(islice(rows, chunk))
                cursor = self._page(ctx, file_id, cursor, page, opened.rows_total)
                _ = self._queue.heartbeat(job, lease=lease)

    def _page(
        self,
        ctx: TenantContext,
        file_id: UUID,
        cursor: int,
        page: list[FileRow],
        rows_total: int,
    ) -> int:
        with tenant_session(ctx, factory=self._session_factory) as s:
            file = self._lock(s, file_id)
            if file.cursor != cursor:
                raise FileCursorMovedError
            if page:
                _ = ingest_page(s, file, page, store=self._store, queue=self._queue)
            else:
                # Fewer rows than counted at upload: nothing is left to read.
                file.cursor = rows_total
            return file.cursor

    def _complete(self, session: Session, file_id: UUID) -> None:
        file = self._lock(session, file_id)
        file.status = FileStatus.COMPLETED
        file.finished_at = self._clock()
        _ = publish_monitoring_event(
            session,
            TRANSACTION_FILE_COMPLETED,
            entity="transaction_file",
            entity_id=file.id,
            key=str(file.id),
            data={
                "file_id": str(file.id),
                "status": file.status.value,
                "rows_accepted": file.rows_accepted,
                "rows_rejected": file.rows_rejected,
            },
        )
        logger.info(
            "transaction file ingest completed",
            extra={
                "file_id": str(file.id),
                "rows_accepted": file.rows_accepted,
                "rows_duplicate": file.rows_duplicate,
                "rows_rejected": file.rows_rejected,
            },
        )
