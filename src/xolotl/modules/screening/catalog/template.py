"""The private-list template: reading a CSV or XLSX upload and validating its rows.

SP-03 spec §7. A file is read row by row (stopping past the row limit), then every
row is checked; any failure rejects the upload with ``PrivateListInvalidError``
carrying each row error located as ``("row", <file row>, <column>)`` (or
``("file",)`` for the file itself) with a fixed message per code. A cell's value
never reaches an error, a message or a log line.

Mapping choices (plan Deviation 5): ``source_ref`` is ``"clave <clave>"`` or empty,
never a row number; ``observaciones`` stays in the sealed upload and never enters an
entry; aliases are ``STRONG`` names; ``otro_identificador`` is an ``OTHER``
identifier; ``fecha_alta`` is ``listed_on``. RFC and CURP are normalized and checked
with the party validators' rules (plan Deviation 7).
"""

import csv
import io
import re
import unicodedata
import zipfile
import zlib
from collections.abc import Callable, Generator, Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from enum import StrEnum

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from xolotl.modules.party.enums import IdentifierType
from xolotl.modules.party.identifiers import IdentifierError, validate
from xolotl.modules.screening.errors import PrivateListInvalidError
from xolotl.modules.screening.lists import (
    EntityType,
    EntryIdentifier,
    EntryName,
    IdType,
    ListEntry,
    NameQuality,
    PartialDate,
)
from xolotl.platform.errors import ErrorItem

TEMPLATE_COLUMNS: tuple[str, ...] = (
    "clave",
    "tipo",
    "nombre_completo",
    "nombres",
    "primer_apellido",
    "segundo_apellido",
    "razon_social",
    "alias",
    "fecha_nacimiento",
    "rfc",
    "curp",
    "nacionalidad",
    "otro_identificador",
    "fecha_alta",
    "observaciones",
)
REQUIRED_COLUMNS: tuple[str, ...] = ("clave", "tipo")
CSV_TYPE = "text/csv"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_REPORTED_ERRORS = 200
# The widest XLSX row read: the template, with room for blank and stray columns.
MAX_XLSX_COLUMNS = 2 * len(TEMPLATE_COLUMNS)


class TemplateError(StrEnum):
    EMPTY_FILE = "EMPTY_FILE"
    UNREADABLE_FILE = "UNREADABLE_FILE"
    MISSING_COLUMN = "MISSING_COLUMN"
    UNKNOWN_COLUMN = "UNKNOWN_COLUMN"
    TOO_MANY_ROWS = "TOO_MANY_ROWS"
    REQUIRED = "REQUIRED"
    INVALID_TYPE = "INVALID_TYPE"
    NAME_AMBIGUOUS = "NAME_AMBIGUOUS"
    INVALID_DATE = "INVALID_DATE"
    INVALID_RFC = "INVALID_RFC"
    INVALID_CURP = "INVALID_CURP"
    RFC_TYPE_MISMATCH = "RFC_TYPE_MISMATCH"
    INVALID_COUNTRY = "INVALID_COUNTRY"
    DUPLICATE_KEY = "DUPLICATE_KEY"


MESSAGES: Mapping[TemplateError, str] = {
    TemplateError.EMPTY_FILE: "the file has no data rows",
    TemplateError.UNREADABLE_FILE: "the file cannot be read as the template",
    TemplateError.MISSING_COLUMN: "a required column (clave or tipo) is missing",
    TemplateError.UNKNOWN_COLUMN: "a column is outside the template or repeated",
    TemplateError.TOO_MANY_ROWS: "the file has more rows than allowed",
    TemplateError.REQUIRED: "a required value is missing",
    TemplateError.INVALID_TYPE: "tipo must be PF or PM",
    TemplateError.NAME_AMBIGUOUS: "the name is given in more than one form",
    TemplateError.INVALID_DATE: "invalid date",
    TemplateError.INVALID_RFC: "invalid RFC",
    TemplateError.INVALID_CURP: "invalid CURP",
    TemplateError.RFC_TYPE_MISMATCH: "the RFC length does not match tipo",
    TemplateError.INVALID_COUNTRY: "invalid country code",
    TemplateError.DUPLICATE_KEY: "the key repeats an earlier row",
}

_PERSON_RFC_LEN = 13
_COMPANY_RFC_LEN = 12
_PARTIAL_DATE = re.compile(r"^(\d{4})(?:-(\d{2})(?:-(\d{2}))?)?$", re.ASCII)
_FULL_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$", re.ASCII)
_COUNTRY = re.compile(r"^[A-Z]{2}$", re.ASCII)
_NOT_ALNUM = re.compile(r"[^0-9A-Z]")
_FILE_LOC: tuple[str | int, ...] = ("file",)
_XLSX_ERRORS = (
    InvalidFileException,
    zipfile.BadZipFile,
    zlib.error,
    # The XML parsers' ParseError (ElementTree's and lxml's) is a SyntaxError.
    SyntaxError,
    EOFError,
    KeyError,
    ValueError,
    OSError,
)


@dataclass(frozen=True, slots=True)
class TemplateRow:
    """One data row: ``number`` is the 1-based file row (the header is row 1).

    ``stray_column`` is the 1-based number of the row's first non-blank cell under
    no template column (past the header or under a blank header cell), if any.
    """

    number: int
    cells: Mapping[str, str]
    stray_column: int | None = None

    def cell(self, column: str) -> str:
        """Give a cell, trimmed, or the empty string when the column is absent.

        Returns:
            The cell's text.

        """
        return self.cells.get(column, "")


def _item(loc: tuple[str | int, ...], code: TemplateError) -> ErrorItem:
    return ErrorItem(loc=loc, msg=MESSAGES[code], type=code.value)


def _file_items(code: TemplateError) -> list[ErrorItem]:
    return [_item(_FILE_LOC, code)]


def _file_error(code: TemplateError) -> PrivateListInvalidError:
    return PrivateListInvalidError(MESSAGES[code], errors=_file_items(code))


# ---- reading -------------------------------------------------------------------------


def _csv_records(data: bytes) -> Generator[list[str]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise _file_error(TemplateError.UNREADABLE_FILE) from None
    first_line = text.split("\n", 1)[0]
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    try:
        yield from csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    except csv.Error:
        raise _file_error(TemplateError.UNREADABLE_FILE) from None


def _xlsx_cell(value: object) -> str:
    # Dates as ISO text; whole numbers without a decimal point (a year, a numeric
    # clave); everything else as its text.
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _xlsx_records(data: bytes) -> Generator[list[str]]:
    # A read-only workbook parses its sheet lazily, so a corrupt sheet fails while
    # rows are read; and it pads each row to the sheet's declared width, which the
    # file controls, so the width read is capped.
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except _XLSX_ERRORS:
        raise _file_error(TemplateError.UNREADABLE_FILE) from None
    try:
        sheets = book.worksheets
        if not sheets:
            raise _file_error(TemplateError.EMPTY_FILE)
        rows = sheets[0].iter_rows(values_only=True, max_col=MAX_XLSX_COLUMNS)
        for row in rows:
            yield [_xlsx_cell(value) for value in row]
    except _XLSX_ERRORS:
        raise _file_error(TemplateError.UNREADABLE_FILE) from None
    finally:
        book.close()


def _header(record: Sequence[str]) -> list[str]:
    names = [cell.strip().lower() for cell in record]
    present = [name for name in names if name]
    if any(column not in present for column in REQUIRED_COLUMNS):
        raise _file_error(TemplateError.MISSING_COLUMN)
    if any(name not in TEMPLATE_COLUMNS for name in present) or len(
        set(present)
    ) != len(present):
        raise _file_error(TemplateError.UNKNOWN_COLUMN)
    return names


def read_template(
    data: bytes, content_type: str, *, max_rows: int
) -> list[TemplateRow]:
    """Read an upload into rows keyed by template column.

    A CSV is UTF-8 (a BOM is tolerated) with a comma or semicolon delimiter chosen
    from the header line; an XLSX is read from its first sheet, dates as ISO text.
    Header names are trimmed and lower-cased; blank header cells and blank rows are
    skipped, and cells are trimmed. Rows are read one at a time and reading stops
    at the first data row past ``max_rows``; an XLSX row is read up to
    ``MAX_XLSX_COLUMNS`` cells.

    Returns:
        The data rows in file order.

    Raises:
        PrivateListInvalidError: When the file cannot be read, has no data row,
            lacks ``clave`` or ``tipo``, has a column outside the template or
            repeated, or more than ``max_rows`` data rows; the one error is located
            at ``("file",)``.

    """
    records = _xlsx_records(data) if content_type == XLSX_TYPE else _csv_records(data)
    names: list[str] | None = None
    rows: list[TemplateRow] = []
    try:
        for number, record in enumerate(records, start=1):
            if not any(cell.strip() for cell in record):
                continue
            if names is None:
                names = _header(record)
            elif len(rows) == max_rows:
                code = TemplateError.TOO_MANY_ROWS
                raise PrivateListInvalidError(MESSAGES[code], errors=_file_items(code))
            else:
                rows.append(_template_row(number, names, record))
    finally:
        records.close()
    if not rows:
        code = TemplateError.EMPTY_FILE
        raise PrivateListInvalidError(MESSAGES[code], errors=_file_items(code))
    return rows


def _template_row(
    number: int, names: Sequence[str], record: Sequence[str]
) -> TemplateRow:
    cells = {
        name: cell.strip() for name, cell in zip(names, record, strict=False) if name
    }
    stray = next(
        (
            column
            for column, cell in enumerate(record, start=1)
            if cell.strip() and (column > len(names) or not names[column - 1])
        ),
        None,
    )
    return TemplateRow(number=number, cells=cells, stray_column=stray)


# ---- validating ----------------------------------------------------------------------


def _clean(text: str) -> str:
    return " ".join(text.split())


def fold_name(text: str) -> str:
    """Fold a name for keying: no accents, case-folded, single spaces.

    Returns:
        The folded name.

    """
    decomposed = unicodedata.normalize("NFKD", text)
    stripped = "".join(c for c in decomposed if not unicodedata.combining(c))
    return _clean(stripped.casefold())


class _Row:
    """The checks of one row; ``errors`` collects what failed."""

    def __init__(self, row: TemplateRow) -> None:
        self.row: TemplateRow = row
        self.errors: list[ErrorItem] = []

    def fail(self, column: str, code: TemplateError) -> None:
        self.errors.append(_item(("row", self.row.number, column), code))

    def cell(self, column: str) -> str:
        return self.row.cell(column)

    def entity_type(self) -> EntityType | None:
        tipo = self.cell("tipo").upper()
        if not tipo:
            self.fail("tipo", TemplateError.REQUIRED)
            return None
        if tipo == "PF":
            return EntityType.PERSON
        if tipo == "PM":
            return EntityType.ORGANIZATION
        self.fail("tipo", TemplateError.INVALID_TYPE)
        return None

    def person_name(self) -> EntryName | None:
        full = _clean(self.cell("nombre_completo"))
        given = _clean(self.cell("nombres"))
        surnames = [
            _clean(self.cell(c)) for c in ("primer_apellido", "segundo_apellido")
        ]
        if self.cell("razon_social"):
            self.fail("razon_social", TemplateError.NAME_AMBIGUOUS)
            return None
        if full and (given or any(surnames)):
            self.fail("nombre_completo", TemplateError.NAME_AMBIGUOUS)
            return None
        if full:
            return EntryName(full=full, quality=NameQuality.PRIMARY)
        if not given and not surnames[0]:
            self.fail("nombre_completo", TemplateError.REQUIRED)
            return None
        if not given or not surnames[0]:
            self.fail(
                "nombres" if not given else "primer_apellido", TemplateError.REQUIRED
            )
            return None
        surname = " ".join(s for s in surnames if s)
        return EntryName(
            full=f"{given} {surname}",
            quality=NameQuality.PRIMARY,
            given=given,
            surname=surname,
        )

    def company_name(self) -> EntryName | None:
        for column in (
            "nombre_completo",
            "nombres",
            "primer_apellido",
            "segundo_apellido",
        ):
            if self.cell(column):
                self.fail(column, TemplateError.NAME_AMBIGUOUS)
                return None
        name = _clean(self.cell("razon_social"))
        if not name:
            self.fail("razon_social", TemplateError.REQUIRED)
            return None
        return EntryName(full=name, quality=NameQuality.PRIMARY)

    def birth_dates(self) -> tuple[PartialDate, ...]:
        text = self.cell("fecha_nacimiento")
        if not text:
            return ()
        match = _PARTIAL_DATE.fullmatch(text)
        if match is None:
            self.fail("fecha_nacimiento", TemplateError.INVALID_DATE)
            return ()
        year, month, day = match.groups()
        try:
            return (
                PartialDate(
                    int(year),
                    None if month is None else int(month),
                    None if day is None else int(day),
                ),
            )
        except ValueError:
            self.fail("fecha_nacimiento", TemplateError.INVALID_DATE)
            return ()

    def listed_on(self) -> date | None:
        text = self.cell("fecha_alta")
        if not text:
            return None
        if _FULL_DATE.fullmatch(text) is not None:
            try:
                return date.fromisoformat(text)
            except ValueError:
                pass
        self.fail("fecha_alta", TemplateError.INVALID_DATE)
        return None

    def identifiers(self, entity_type: EntityType | None) -> list[EntryIdentifier]:
        found: list[EntryIdentifier] = []
        rfc = self._checked("rfc", IdentifierType.RFC, TemplateError.INVALID_RFC)
        if rfc is not None:
            expected = (
                _COMPANY_RFC_LEN
                if entity_type is EntityType.ORGANIZATION
                else _PERSON_RFC_LEN
            )
            if entity_type is not None and len(rfc) != expected:
                self.fail("rfc", TemplateError.RFC_TYPE_MISMATCH)
            else:
                found.append(EntryIdentifier(IdType.RFC, rfc, rfc, country="MX"))
        curp = self._checked("curp", IdentifierType.CURP, TemplateError.INVALID_CURP)
        if curp is not None:
            found.append(EntryIdentifier(IdType.CURP, curp, curp, country="MX"))
        other = self.cell("otro_identificador")
        normalized = _NOT_ALNUM.sub("", unicodedata.normalize("NFKD", other.upper()))
        if normalized:
            found.append(EntryIdentifier(IdType.OTHER, other, normalized))
        return found

    def _checked(
        self, column: str, id_type: IdentifierType, code: TemplateError
    ) -> str | None:
        raw = self.cell(column)
        if not raw:
            return None
        result = validate(id_type, raw)
        if isinstance(result, IdentifierError):
            self.fail(column, code)
            return None
        return result.value

    def nationalities(self) -> tuple[str, ...]:
        codes = [c.strip().upper() for c in self.cell("nacionalidad").split("|")]
        codes = [c for c in codes if c]
        if any(_COUNTRY.fullmatch(c) is None for c in codes):
            self.fail("nacionalidad", TemplateError.INVALID_COUNTRY)
            return ()
        return tuple(codes)

    def aliases(self) -> tuple[EntryName, ...]:
        names = (_clean(a) for a in self.cell("alias").split("|"))
        return tuple(EntryName(full=n, quality=NameQuality.STRONG) for n in names if n)


def _entry(
    checks: _Row, *, list_code: str, key_tag: Callable[[str], str]
) -> ListEntry | None:
    entity_type = checks.entity_type()
    name: EntryName | None = None
    if entity_type is EntityType.PERSON:
        name = checks.person_name()
    elif entity_type is EntityType.ORGANIZATION:
        name = checks.company_name()
    birth_dates = checks.birth_dates()
    identifiers = checks.identifiers(entity_type)
    nationalities = checks.nationalities()
    listed_on = checks.listed_on()
    aliases = checks.aliases()
    if checks.errors or entity_type is None or name is None:
        return None
    clave = checks.cell("clave")
    key = clave or key_tag(f"{fold_name(name.full)}|{checks.cell('fecha_nacimiento')}")
    return ListEntry(
        list_code=list_code,
        entry_key=key,
        entity_type=entity_type,
        primary_name=name,
        aliases=aliases,
        birth_dates=birth_dates,
        nationalities=nationalities,
        identifiers=tuple(identifiers),
        listed_on=listed_on,
        source_ref=f"clave {clave}" if clave else "",
    )


def rows_to_entries(
    rows: Iterable[TemplateRow], *, list_code: str, key_tag: Callable[[str], str]
) -> list[ListEntry]:
    """Validate every row and map it to a list entry.

    ``key_tag`` turns the folded name and birth date of a row without ``clave``
    into its entry key.

    Returns:
        One entry per row, in file order.

    Raises:
        PrivateListInvalidError: When any row fails; it carries the first
            ``MAX_REPORTED_ERRORS`` row errors, and no entry is returned.

    """
    entries: list[ListEntry] = []
    errors: list[ErrorItem] = []
    seen: set[str] = set()
    for row in rows:
        checks = _Row(row)
        if row.stray_column is not None:
            # A value outside the template would be silently dropped; it is
            # located by its column number, as it has no template name.
            checks.errors.append(
                _item(
                    ("row", row.number, row.stray_column), TemplateError.UNKNOWN_COLUMN
                )
            )
        # A repeated clave is refused even when its first row failed; a keyed tag
        # is known only once the row's name is.
        clave = row.cell("clave")
        if clave:
            if clave in seen:
                checks.fail("clave", TemplateError.DUPLICATE_KEY)
            seen.add(clave)
        entry = _entry(checks, list_code=list_code, key_tag=key_tag)
        if entry is not None and not clave:
            if entry.entry_key in seen:
                checks.fail("clave", TemplateError.DUPLICATE_KEY)
                entry = None
            else:
                seen.add(entry.entry_key)
        if entry is not None:
            entries.append(entry)
        errors.extend(checks.errors)
    if errors:
        detail = f"{len(errors)} errors in the file"
        raise PrivateListInvalidError(detail, errors=errors[:MAX_REPORTED_ERRORS])
    return entries
